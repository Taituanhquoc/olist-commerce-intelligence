"""Published outputs (dashboard, Excel, README, notebooks) must agree with facts.json."""
import re
from pathlib import Path

import nbformat
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_dashboard_totals_match_facts(facts):
    d, k = facts["dashboard"], facts["analysis"]["kpi"]
    assert d["oc_revenue"] == pytest.approx(k["revenue"], abs=1.0)
    assert d["ic_revenue"] == pytest.approx(k["revenue"], abs=1.0)
    assert d["oc_orders"] == k["orders"]
    assert d["oc_late"] == k["late_orders"]
    assert d["oc_delivered"] == k["delivered_orders"]


def test_dashboard_is_self_contained():
    html = (ROOT / "dashboard" / "index.html").read_text(encoding="utf-8")
    assert "<script src=" not in html, "dashboard must not load external scripts"
    # Portfolio extension owns GitHub Pages; the original complete dashboard remains available separately.
    assert (ROOT / "docs" / "index.html").read_bytes() == (ROOT / "promise_lab" / "index.html").read_bytes()
    assert (ROOT / "docs" / "promise_lab.html").read_bytes() == (ROOT / "promise_lab" / "index.html").read_bytes()


def test_excel_kpis_are_formulas_and_match(facts):
    ws = load_workbook(ROOT / "excel" / "olist_ecommerce_kpis.xlsx")["Summary"]
    for row in range(5, 16):
        assert str(ws.cell(row, 2).value).startswith("="), f"Summary!B{row} should be a formula"
    if facts["excel_check"]["recalculated_by_excel"]:
        assert facts["excel_check"]["mismatches"] == 0


def test_readme_generated_and_figures_exist():
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    assert "{{" not in text and "TODO" not in text
    for path in re.findall(r"\]\((reports/figures/[^)]+)\)", text):
        assert (ROOT / path).exists(), path


@pytest.mark.parametrize("name", ["01_audit", "02_cleaning", "03_eda", "04_analysis"])
def test_notebooks_executed_without_errors(name):
    nb = nbformat.read(ROOT / "notebooks" / f"{name}.ipynb", as_version=4)
    code = [c for c in nb.cells if c.cell_type == "code"]
    assert code and all(c.execution_count for c in code)
    assert not any(o.get("output_type") == "error" for c in code for o in c.outputs)
