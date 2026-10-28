"""Build excel/olist_ecommerce_kpis.xlsx: formula-driven KPI workbook with an impact model.

Data sheets hold aggregated values from the cleaned tables; every KPI, share, rank, growth
rate and impact figure is an Excel formula on top of them, so the workbook recalculates if a
reviewer changes an input. src/excel_verify.py recalculates it and checks it against facts.json.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src import facts
from src.clean import load
from src.config import EXCEL_PATH

NAVY, BLUE, GREY = "1F2A44", "2A78D6", "F0EFEC"
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
BASE_FILL = PatternFill("solid", fgColor="E8F1FC")
THIN = Border(bottom=Side(style="thin", color="C3C2B7"))
BRL, BRL2, PCT, PCT2, INT, DEC2 = 'R$ #,##0', 'R$ #,##0.00', '0.0%', '0.00%', '#,##0', '0.00'


def header(ws, row: int, cols: list[str]) -> None:
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, w: dict[str, int]) -> None:
    for col, width in w.items():
        ws.column_dimensions[col].width = width


def title(ws, text: str, sub: str) -> None:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(italic=True, size=9, color="52514E")


def build() -> None:
    orders, items = load()
    f = facts.read()["analysis"]
    wb = Workbook()

    # ---------------------------------------------------------------- Monthly
    ws = wb.active
    ws.title = "Monthly"
    title(ws, "Monthly KPIs", "Values: revenue, orders, customers from the cleaned order table. Columns F-I are formulas.")
    m = orders.groupby("purchase_month").agg(revenue=("revenue", "sum"), orders=("order_id", "size"),
                                             customers=("customer_unique_id", "nunique"),
                                             new=("customer_order_seq", lambda s: int((s == 1).sum()))).reset_index()
    header(ws, 4, ["Month", "Revenue", "Orders", "Customers", "New customers", "AOV", "MoM %", "YoY %", "Running revenue"])
    first, last = 5, 5 + len(m) - 1
    for i, r in enumerate(m.itertuples(), first):
        ws.cell(i, 1, r.purchase_month.strftime("%Y-%m"))
        ws.cell(i, 2, round(r.revenue, 2)).number_format = BRL
        ws.cell(i, 3, r.orders).number_format = INT
        ws.cell(i, 4, r.customers).number_format = INT
        ws.cell(i, 5, r.new).number_format = INT
        ws.cell(i, 6, f"=B{i}/C{i}").number_format = BRL2
        ws.cell(i, 7, f"=B{i}/B{i-1}-1" if i > first else None).number_format = PCT
        ws.cell(i, 8, f"=B{i}/B{i-12}-1" if i - 12 >= first else None).number_format = PCT
        ws.cell(i, 9, f"=SUM($B${first}:B{i})").number_format = BRL
    ws.conditional_formatting.add(f"G{first + 1}:H{last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="D03B3B")))
    ws.conditional_formatting.add(f"G{first + 1}:H{last}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color="006300")))
    ws.conditional_formatting.add(f"B{first}:B{last}", DataBarRule(start_type="min", end_type="max", color=BLUE))
    widths(ws, {"A": 10, "B": 14, "C": 10, "D": 11, "E": 13, "F": 11, "G": 9, "H": 9, "I": 16})
    ch = LineChart()
    ch.title, ch.y_axis.title, ch.height, ch.width = "Monthly revenue (R$)", "Revenue", 8, 18
    ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ws.add_chart(ch, "K4")
    ws.freeze_panes = "A5"
    monthly_rng = (first, last)

    # ---------------------------------------------------------------- Category
    wc = wb.create_sheet("Category")
    title(wc, "Category performance", "Values: item revenue, items, freight by category. Share, rank, cumulative share and freight ratio are formulas.")
    cat = (items.groupby("category").agg(revenue=("price", "sum"), items=("order_item_id", "size"), freight=("freight_value", "sum"))
           .sort_values("revenue", ascending=False).reset_index())
    header(wc, 4, ["Category", "Revenue", "Items", "Freight", "Revenue share", "Rank", "Cumulative share", "Freight / revenue"])
    c1, cl = 5, 5 + len(cat) - 1
    for i, r in enumerate(cat.itertuples(), c1):
        wc.cell(i, 1, r.category)
        wc.cell(i, 2, round(r.revenue, 2)).number_format = BRL
        wc.cell(i, 3, r.items).number_format = INT
        wc.cell(i, 4, round(r.freight, 2)).number_format = BRL
        wc.cell(i, 5, f"=B{i}/SUM($B${c1}:$B${cl})").number_format = PCT2
        wc.cell(i, 6, f"=RANK(B{i},$B${c1}:$B${cl})")
        wc.cell(i, 7, f"=SUM($B${c1}:B{i})/SUM($B${c1}:$B${cl})").number_format = PCT
        wc.cell(i, 8, f"=D{i}/B{i}").number_format = PCT
    wc.conditional_formatting.add(f"H{c1}:H{cl}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="EB6834"))
    wc.conditional_formatting.add(f"B{c1}:B{cl}", DataBarRule(start_type="min", end_type="max", color=BLUE))
    widths(wc, {"A": 34, "B": 14, "C": 9, "D": 13, "E": 12, "F": 7, "G": 13, "H": 13})
    bc = BarChart()
    bc.type, bc.title, bc.height, bc.width = "bar", "Top 15 categories by revenue", 10, 16
    bc.add_data(Reference(wc, min_col=2, min_row=4, max_row=c1 + 14), titles_from_data=True)
    bc.set_categories(Reference(wc, min_col=1, min_row=c1, max_row=c1 + 14))
    bc.y_axis.majorGridlines = None
    wc.add_chart(bc, "J4")
    wc.freeze_panes = "A5"

    # ---------------------------------------------------------------- Region (state)
    wr = wb.create_sheet("Region")
    title(wr, "Logistics and reviews by customer state",
          "Values: order, delivery and review sums by state. Rates and averages are formulas; F2 holds the national late rate.")
    d = orders.copy()
    d["late"] = d.is_late.fillna(False).astype(bool) & d.is_delivered
    st = d.groupby("customer_state").agg(
        orders=("order_id", "size"), revenue=("revenue", "sum"), delivered=("is_delivered", "sum"),
        late=("late", "sum"), days=("delivery_days", "sum"), rsum=("review_score", "sum"),
        rcount=("review_score", "count")).sort_values("orders", ascending=False).reset_index()
    header(wr, 4, ["State", "Orders", "Revenue", "Delivered", "Late", "Delivery-day sum", "Review sum", "Reviews",
                   "Late rate", "Avg delivery days", "Avg review", "vs national"])
    r1, rl = 5, 5 + len(st) - 1
    wr["E2"], wr["F2"] = "National late rate:", f"=SUM(E{r1}:E{rl})/SUM(D{r1}:D{rl})"
    wr["F2"].number_format = PCT
    for i, r in enumerate(st.itertuples(), r1):
        wr.cell(i, 1, r.customer_state)
        wr.cell(i, 2, int(r.orders)).number_format = INT
        wr.cell(i, 3, round(r.revenue, 2)).number_format = BRL
        wr.cell(i, 4, int(r.delivered)).number_format = INT
        wr.cell(i, 5, int(r.late)).number_format = INT
        wr.cell(i, 6, round(r.days, 4)).number_format = INT
        wr.cell(i, 7, int(r.rsum)).number_format = INT
        wr.cell(i, 8, int(r.rcount)).number_format = INT
        wr.cell(i, 9, f"=IF(D{i}=0,\"\",E{i}/D{i})").number_format = PCT
        wr.cell(i, 10, f"=IF(D{i}=0,\"\",F{i}/D{i})").number_format = "0.0"
        wr.cell(i, 11, f"=IF(H{i}=0,\"\",G{i}/H{i})").number_format = DEC2
        wr.cell(i, 12, f"=IF(AND(D{i}>=300,I{i}>$F$2),\"Above national\",\"\")")
    wr.conditional_formatting.add(f"I{r1}:I{rl}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="D03B3B"))
    wr.conditional_formatting.add(f"K{r1}:K{rl}", ColorScaleRule(start_type="min", start_color="D03B3B", mid_type="percentile",
                                                                 mid_value=50, mid_color="FFFFFF", end_type="max", end_color="2A78D6"))
    widths(wr, {"A": 7, "B": 9, "C": 14, "D": 10, "E": 8, "F": 14, "G": 11, "H": 9, "I": 9, "J": 11, "K": 10, "L": 15})
    wr.freeze_panes = "A5"

    # ---------------------------------------------------------------- RFM
    wf = wb.create_sheet("RFM")
    title(wf, "RFM segments", "Values: customers and revenue per segment. Shares and averages are formulas.")
    seg = pd.DataFrame(f["rfm"]["segments"]).T.reset_index().rename(columns={"index": "segment"})
    header(wf, 4, ["Segment", "Customers", "Revenue", "Customer share", "Revenue share", "Revenue per customer", "Repeat buyers?"])
    s1, sl = 5, 5 + len(seg) - 1
    for i, r in enumerate(seg.itertuples(), s1):
        wf.cell(i, 1, r.segment)
        wf.cell(i, 2, int(r.customers)).number_format = INT
        wf.cell(i, 3, round(float(r.revenue), 2)).number_format = BRL
        wf.cell(i, 4, f"=B{i}/SUM($B${s1}:$B${sl})").number_format = PCT
        wf.cell(i, 5, f"=C{i}/SUM($C${s1}:$C${sl})").number_format = PCT
        wf.cell(i, 6, f"=C{i}/B{i}").number_format = BRL2
        wf.cell(i, 7, "Yes" if r.segment in ("Champions", "Loyal") else "No")
    widths(wf, {"A": 16, "B": 11, "C": 14, "D": 14, "E": 13, "F": 18, "G": 14})
    bf = BarChart()
    bf.title, bf.height, bf.width = "Customer share vs revenue share", 8, 16
    bf.add_data(Reference(wf, min_col=4, max_col=5, min_row=4, max_row=sl), titles_from_data=True)
    bf.set_categories(Reference(wf, min_col=1, min_row=s1, max_row=sl))
    wf.add_chart(bf, "I4")

    # ---------------------------------------------------------------- Payments
    wp = wb.create_sheet("Payments")
    title(wp, "Payment mix", "Values: orders and revenue by primary payment type. Share and AOV are formulas.")
    pm = orders.groupby(orders.primary_payment_type.fillna("unknown")).agg(orders=("order_id", "size"), revenue=("revenue", "sum")) \
        .sort_values("orders", ascending=False).reset_index()
    header(wp, 4, ["Payment type", "Orders", "Revenue", "Order share", "AOV"])
    p1, pl = 5, 5 + len(pm) - 1
    for i, r in enumerate(pm.itertuples(), p1):
        wp.cell(i, 1, r.primary_payment_type)
        wp.cell(i, 2, int(r.orders)).number_format = INT
        wp.cell(i, 3, round(r.revenue, 2)).number_format = BRL
        wp.cell(i, 4, f"=B{i}/SUM($B${p1}:$B${pl})").number_format = PCT
        wp.cell(i, 5, f"=C{i}/B{i}").number_format = BRL2
    widths(wp, {"A": 14, "B": 10, "C": 14, "D": 12, "E": 11})

    # ---------------------------------------------------------------- Impact model
    wi = wb.create_sheet("Impact model")
    title(wi, "Growth actions: annual revenue impact", "Yellow cells are inputs (change them). Blue cells are base values from the analysis.")
    imp = f["impact"]
    rows = [
        ("Base values", None, None),
        ("Annual new customers (2018 run-rate)", imp["annual_new_customers"], INT),
        ("Average repeat-order revenue (R$)", imp["avg_repeat_order_revenue"], BRL2),
        ("Annual orders (2018 run-rate)", imp["annual_orders"], INT),
        ("Average add-on revenue per multi-item order (R$)", imp["a3_avg_addon_revenue"], BRL2),
        ("Late orders avoided per year (target states to national rate)", imp["a2_late_orders_avoided"], INT),
        ("Repeat-rate gap, on-time vs late first order (pp)", f["retention_vs_delivery"]["gap_pp"], DEC2),
        ("Annualised 2018 revenue (R$)", f["trend"]["annualised_revenue_2018"], BRL),
        ("Assumptions", None, None),
        ("Repeat-rate lift from second-purchase programme (pp)", f["assumptions"]["repeat_lift_pp"], DEC2),
        ("Multi-item order share lift from cross-sell (pp)", f["assumptions"]["basket_lift_pp"], DEC2),
    ]
    r = 4
    ref = {}
    for label, val, fmt in rows:
        wi.cell(r, 1, label)
        if val is None:
            wi.cell(r, 1).font = Font(bold=True, color=NAVY)
        else:
            c = wi.cell(r, 2, round(float(val), 4))
            c.number_format = fmt
            c.fill = INPUT_FILL if r >= 13 else BASE_FILL
            ref[label] = f"B{r}"
        r += 1
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10", showErrorMessage=True,
                        error="Enter a lift between 0 and 10 percentage points")
    wi.add_data_validation(dv)
    dv.add("B13:B14")
    out = r + 1
    wi.cell(out, 1, "Results").font = Font(bold=True, color=NAVY)
    res = [
        ("A1 Second-purchase programme (R$/yr)", f"=B5*B13/100*B6"),
        ("A2 Delivery fix in high-late states (R$/yr, via repeat rate)", f"=B9*B10/100*B6"),
        ("A3 Cross-sell to build baskets (R$/yr)", f"=B7*B14/100*B8"),
        ("Total (R$/yr)", f"=SUM(B{out + 1}:B{out + 3})"),
        ("Total as % of annualised 2018 revenue", f"=B{out + 4}/B11"),
    ]
    for k, (label, formula) in enumerate(res, out + 1):
        wi.cell(k, 1, label)
        wi.cell(k, 2, formula).number_format = PCT2 if "%" in label else BRL
        wi.cell(k, 2).font = Font(bold=True)
    sens = out + 8
    wi.cell(sens, 1, "Sensitivity: A1 revenue by repeat-rate lift").font = Font(bold=True, color=NAVY)
    header(wi, sens + 1, ["Repeat-rate lift (pp)", "A1 revenue (R$/yr)"])
    for k, lift in enumerate([0.5, 1.0, 1.5, 2.0, 2.5, 3.0], sens + 2):
        wi.cell(k, 1, lift).number_format = DEC2
        wi.cell(k, 2, f"=$B$5*A{k}/100*$B$6").number_format = BRL
    sc = BarChart()
    sc.title, sc.height, sc.width = "A1 revenue by repeat-rate lift", 7, 12
    sc.add_data(Reference(wi, min_col=2, min_row=sens + 1, max_row=sens + 7), titles_from_data=True)
    sc.set_categories(Reference(wi, min_col=1, min_row=sens + 2, max_row=sens + 7))
    wi.add_chart(sc, "D4")
    widths(wi, {"A": 58, "B": 16})

    # ---------------------------------------------------------------- Summary (first sheet)
    wsu = wb.create_sheet("Summary", 0)
    title(wsu, "Olist e-commerce: executive KPIs (Jan 2017 - Aug 2018)",
          "Every value below is a formula over the data sheets. Source: Olist public dataset (Kaggle), real data.")
    a, b = monthly_rng
    kpis = [
        ("Revenue (item prices, R$)", f"=SUM(Monthly!B{a}:B{b})", BRL),
        ("Orders", f"=SUM(Monthly!C{a}:C{b})", INT),
        ("Average order value (R$)", "=B5/B6", BRL2),
        ("Customers", f"=SUM(RFM!B{s1}:B{sl})", INT),
        ("Repeat purchase rate", f"=SUMIF(RFM!G{s1}:G{sl},\"Yes\",RFM!B{s1}:B{sl})/B8", PCT2),
        ("On-time delivery rate", f"=1-SUM(Region!E{r1}:E{rl})/SUM(Region!D{r1}:D{rl})", PCT),
        ("Average delivery days", f"=SUM(Region!F{r1}:F{rl})/SUM(Region!D{r1}:D{rl})", "0.0"),
        ("Average review score", f"=SUM(Region!G{r1}:G{rl})/SUM(Region!H{r1}:H{rl})", DEC2),
        ("Jan-Aug YoY revenue growth", f"=SUM(Monthly!B{a + 12}:B{b})/SUM(Monthly!B{a}:B{a + 7})-1", PCT),
        ("Top-10 category revenue share", f"=SUMIF(Category!F{c1}:F{cl},\"<=10\",Category!E{c1}:E{cl})", PCT),
        ("Growth actions: total annual impact (R$)", f"='Impact model'!B{out + 4}", BRL),
    ]
    header(wsu, 4, ["KPI", "Value"])
    for k, (label, formula, fmt) in enumerate(kpis, 5):
        wsu.cell(k, 1, label).border = THIN
        c = wsu.cell(k, 2, formula)
        c.number_format, c.font, c.border = fmt, Font(bold=True, size=12), THIN
    wsu.cell(5 + len(kpis) + 1, 1, "Sheets: Monthly, Category, Region, RFM, Payments, Impact model. "
                                   "Rates and averages are recomputed from sums, never typed in.").font = Font(italic=True, size=9)
    widths(wsu, {"A": 44, "B": 18})
    wsu.sheet_view.showGridLines = False

    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(EXCEL_PATH)
    print(f"excel: wrote {EXCEL_PATH.name} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    build()
