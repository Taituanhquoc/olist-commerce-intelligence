"""Recalculate the workbook and check its formula results against facts.json.

On Windows with Microsoft Excel installed, Excel recalculates and saves the file (so cached
values are stored for viewers). Elsewhere (CI, Linux) recalculation is skipped and the check
falls back to confirming that the KPI cells hold formulas; the Python values in facts.json
remain the reference.
"""
from __future__ import annotations

import sys

from openpyxl import load_workbook

from src import facts
from src.config import EXCEL_PATH

# Summary sheet cell -> (facts path, tolerance)
CHECKS = {
    "B5": (("kpi", "revenue"), 0.01), "B6": (("kpi", "orders"), 0), "B7": (("kpi", "aov"), 0.01),
    "B8": (("kpi", "customers"), 0), "B9": (("kpi", "repeat_rate"), 1e-6), "B10": (("kpi", "on_time_rate"), 1e-6),
    "B11": (("kpi", "avg_delivery_days"), 1e-3), "B12": (("kpi", "avg_review"), 1e-5),
    "B13": (("trend", "yoy_revenue_growth"), 1e-4), "B14": (("category", "top10_share"), 1e-5),
    "B15": (("impact", "total_revenue"), 30.0),  # inputs are rounded to 4 dp
}


def recalc_with_excel() -> bool:
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return False
    try:
        xl = win32com.client.DispatchEx("Excel.Application")
    except Exception:  # noqa: BLE001 - Excel not installed
        return False
    xl.Visible, xl.DisplayAlerts = False, False
    try:
        wb = xl.Workbooks.Open(str(EXCEL_PATH.resolve()))
        xl.CalculateFull()
        wb.Save()
        wb.Close(SaveChanges=False)
    finally:
        xl.Quit()
    return True


def main() -> int:
    recalculated = recalc_with_excel()
    an = facts.read()["analysis"]
    formulas = load_workbook(EXCEL_PATH)["Summary"]
    values = load_workbook(EXCEL_PATH, data_only=True)["Summary"]
    results, bad = {}, []
    for cell, ((sec, key), tol) in CHECKS.items():
        if not str(formulas[cell].value).startswith("="):
            bad.append(f"{cell} is not a formula")
            continue
        expected = an[sec][key]
        got = values[cell].value
        results[cell] = {"facts": expected, "excel": got}
        if recalculated and (got is None or abs(float(got) - float(expected)) > max(tol, 1e-9 * abs(expected))):
            bad.append(f"{cell}: Excel {got} vs facts {expected}")
    facts.update("excel_check", {"recalculated_by_excel": recalculated, "cells_checked": len(CHECKS),
                                 "mismatches": len(bad), "revenue_excel": results.get("B5", {}).get("excel")})
    if bad:
        print("excel_verify: FAIL\n  " + "\n  ".join(bad))
        return 1
    mode = "recalculated in Excel and matched facts.json" if recalculated else "formulas present (no Excel to recalculate)"
    print(f"excel_verify: OK - {len(CHECKS)} KPI cells {mode}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
